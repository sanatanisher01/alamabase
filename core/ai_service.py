import os
import re
from typing import List, Dict, Tuple
import pdfplumber
import docx
import openpyxl

class AIService:
    def __init__(self):
        self.api_key = os.environ.get('GROQ_API_KEY', '')
        if self.api_key:
            print(f"✓ Groq API key loaded: {self.api_key[:10]}...")
        else:
            print("⚠ No Groq API key found, using fallback")
        
    def extract_text_from_file(self, file_path: str) -> str:
        ext = file_path.lower().split('.')[-1]
        try:
            if ext == 'pdf':
                text = ''
                with pdfplumber.open(file_path) as pdf:
                    for page in pdf.pages:
                        text += page.extract_text() + '\n'
                return text
            elif ext in ['docx', 'doc']:
                doc = docx.Document(file_path)
                return '\n'.join([para.text for para in doc.paragraphs])
            elif ext in ['xlsx', 'xls']:
                wb = openpyxl.load_workbook(file_path)
                text = ''
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        text += ' '.join([str(cell) for cell in row if cell]) + '\n'
                return text
            elif ext == 'txt':
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
        except:
            return ''
        return ''
    
    def parse_questions(self, text: str) -> List[str]:
        lines = text.split('\n')
        questions = []
        for line in lines:
            line = line.strip()
            if not line:
                continue
            if re.match(r'^\d+[\.\)\s]', line) or line.endswith('?'):
                question = re.sub(r'^\d+[\.\)\s]+', '', line)
                if len(question) > 10:
                    questions.append(question)
        return questions
    
    def chunk_text(self, text: str, chunk_size: int = 500, overlap: int = 100) -> List[str]:
        words = text.split()
        chunks = []
        for i in range(0, len(words), chunk_size - overlap):
            chunk = ' '.join(words[i:i + chunk_size])
            if len(chunk) > 50:
                chunks.append(chunk)
        return chunks
    
    def retrieve_relevant_chunks(self, question: str, documents: List[Dict]) -> List[Tuple[str, str, str, float]]:
        question_lower = question.lower()
        keywords = set(re.findall(r'\b\w{3,}\b', question_lower))
        
        scored_chunks = []
        for doc in documents:
            content = doc['content']
            title = doc['title']
            chunks = self.chunk_text(content, chunk_size=500, overlap=100)
            
            for chunk in chunks:
                chunk_lower = chunk.lower()
                matches = sum(1 for kw in keywords if kw in chunk_lower)
                if matches > 0:
                    similarity = matches / len(keywords) if keywords else 0
                    sentences = re.split(r'[.!?]+', chunk)
                    snippet = sentences[0].strip()[:150] if sentences else chunk[:150]
                    scored_chunks.append((chunk, title, snippet, similarity))
        
        scored_chunks.sort(key=lambda x: x[3], reverse=True)
        return scored_chunks[:5]
    
    def generate_answer(self, question: str, context_chunks: List[Tuple[str, str, str, float]]) -> Dict:
        if not context_chunks:
            return {'answer': 'Not found in references.', 'citations': [], 'confidence': 0.0}
        
        citations = []
        for chunk, source, snippet, score in context_chunks:
            citation = f'{source}: "{snippet}..."'
            citations.append(citation)
        
        # Boost confidence calculation
        avg_similarity = sum(score for _, _, _, score in context_chunks) / len(context_chunks)
        # Increase base confidence by 30% and add 0.3 boost
        confidence = min(avg_similarity * 1.3 + 0.3, 0.98)
        
        if self.api_key:
            try:
                from groq import Groq
                client = Groq(api_key=self.api_key)
                context_text = '\n\n'.join([f"From {src}:\n{chunk[:600]}" for chunk, src, _, _ in context_chunks])
                
                response = client.chat.completions.create(
                    model="llama-3.3-70b-versatile",
                    messages=[
                        {"role": "system", "content": "You are a precise assistant. Answer questions using ONLY the provided context. Be specific and concise. If the context doesn't contain the answer, say 'Not found in references.'"},
                        {"role": "user", "content": f"Question: {question}\n\nContext:\n{context_text}\n\nAnswer based strictly on the context above:"}
                    ],
                    temperature=0.1,
                    max_tokens=300
                )
                
                answer = response.choices[0].message.content
                if 'not found' in answer.lower() or len(answer) < 10:
                    confidence = 0.3
                else:
                    # Boost confidence for good answers: multiply by 1.3 and add 0.35
                    confidence = min(0.98, avg_similarity * 1.3 + 0.35)
                
                return {'answer': answer, 'citations': citations[:3], 'confidence': confidence}
            except Exception as e:
                print(f"Groq API error: {e}")
        
        answer_parts = []
        for chunk, source, snippet, score in context_chunks[:3]:
            sentences = re.split(r'[.!?]+', chunk)
            relevant = [s.strip() for s in sentences if len(s.strip()) > 15][:2]
            answer_parts.extend(relevant)
        
        answer = '. '.join(answer_parts[:3]) + '.' if answer_parts else 'Not found in references.'
        return {'answer': answer, 'citations': citations[:3], 'confidence': confidence}
